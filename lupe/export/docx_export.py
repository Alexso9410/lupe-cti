"""DOCX export formatter for case history and individual IOCs."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document


def _add_ioc_added_section(doc: Document, events: list[dict]) -> None:
    """Add IOC addition event details."""
    additions = [e for e in events if e.get("event_type") == "ioc_added"]
    for event in additions:
        ioc_type = event.get("ioc_type", "unknown")
        ts = event.get("timestamp", "")
        notes = event.get("detail", "")
        text = f"[{ts}] IOC added — Type: {ioc_type}"
        if notes:
            text += f" — Notes: {notes}"
        doc.add_paragraph(text)


def _add_enrichment_table(doc: Document, events: list[dict]) -> None:
    """Add a formatted table of enrichment results."""
    enrichments = [e for e in events if e.get("event_type") == "enrichment"]
    if not enrichments:
        return

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Source"
    hdr[1].text = "Severity"
    hdr[2].text = "Summary"
    hdr[3].text = "Timestamp"

    for event in enrichments:
        row = table.add_row().cells
        row[0].text = event.get("source", "")
        row[1].text = event.get("detail", event.get("severity", ""))
        row[2].text = event.get("summary", event.get("description", ""))
        row[3].text = event.get("timestamp", "")


def _add_analysis_section(doc: Document, events: list[dict]) -> None:
    """Add AI analysis paragraphs."""
    analyses = [e for e in events if e.get("event_type") == "analysis"]
    if not analyses:
        return

    doc.add_heading("AI Analysis", level=3)
    for event in analyses:
        model = event.get("model", "unknown")
        summary = event.get("summary", "")
        timestamp = event.get("timestamp", "")
        doc.add_paragraph(f"[{timestamp}] Model: {model}")
        doc.add_paragraph(summary)


def export_case_to_docx(history: dict, output_path: Path) -> Path:
    """Export a full case history to DOCX format.

    Args:
        history: CaseHistory dict from build_case_history().
        output_path: Path where the .docx file will be written.

    Returns:
        Path to the written file.
    """
    doc = Document()
    case = history.get("case", {})
    stats = history.get("stats", {})
    iocs = history.get("iocs", [])
    notes = history.get("case_notes", [])

    # --- Title ---
    doc.add_heading(f"Case Report: {case.get('name', 'N/A')}", level=1)

    # --- Executive Summary ---
    doc.add_heading("Executive Summary", level=2)
    doc.add_paragraph(f"Description: {case.get('description', '')}")
    doc.add_paragraph(f"Status: {case.get('status', 'unknown')}")
    doc.add_paragraph(f"Created: {case.get('created_at', 'N/A')}")
    doc.add_paragraph(f"IOCs: {stats.get('ioc_count', 0)}")
    doc.add_paragraph(f"Enrichments: {stats.get('enrichment_count', 0)}")
    doc.add_paragraph(f"Analyses: {stats.get('analysis_count', 0)}")
    doc.add_paragraph(f"Top Severity: {stats.get('top_severity', 'info')}")
    doc.add_paragraph(
        f"Exported: {datetime.now(tz=timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"
    )

    if not iocs:
        doc.add_paragraph("No IOCs in case.")
        doc.save(str(output_path))
        return output_path

    # --- Case Notes ---
    if notes:
        doc.add_heading("Case Notes", level=2)
        for note in notes:
            doc.add_paragraph(
                f"[{note.get('created_at', '')}] {note.get('content', '')}"
            )

    # --- IOC Details ---
    for ioc_group in iocs:
        ioc = ioc_group.get("ioc", {})
        events = ioc_group.get("events", [])

        doc.add_heading(
            f"[{ioc.get('type', 'unknown')}] {ioc.get('value', 'N/A')}", level=2
        )

        # IOC addition info
        _add_ioc_added_section(doc, events)

        # Enrichment table
        _add_enrichment_table(doc, events)

        # AI Analysis
        _add_analysis_section(doc, events)

    doc.save(str(output_path))
    return output_path


def export_ioc_to_docx(
    ioc_data: dict, case_info: dict | None, output_path: Path
) -> Path:
    """Export a single IOC's data to DOCX format.

    Args:
        ioc_data: Single IOC group dict with 'ioc' and 'events' keys.
        case_info: Optional case metadata dict.
        output_path: Path where the .docx file will be written.

    Returns:
        Path to the written file.
    """
    doc = Document()
    ioc = ioc_data.get("ioc", {})
    events = ioc_data.get("events", [])

    if case_info:
        doc.add_heading(f"Case: {case_info.get('name', 'N/A')}", level=1)

    doc.add_heading(
        f"[{ioc.get('type', 'unknown')}] {ioc.get('value', 'N/A')}", level=2
    )

    _add_enrichment_table(doc, events)
    _add_analysis_section(doc, events)

    doc.save(str(output_path))
    return output_path
