"""Tests for DOCX export formatter."""

from __future__ import annotations

import pytest


@pytest.fixture
def sample_history():
    """A minimal CaseHistory dict for testing."""
    return {
        "case": {
            "id": 1,
            "name": "Incident-2026-001",
            "description": "Phishing investigation",
            "status": "open",
            "created_at": "2026-07-06 10:00:00",
        },
        "iocs": [
            {
                "ioc": {
                    "id": 1,
                    "type": "ipv4",
                    "value": "8.8.8.8",
                    "added_at": "2026-07-06 10:00:00",
                },
                "events": [
                    {
                        "event_type": "enrichment",
                        "timestamp": "2026-07-06 10:01:00",
                        "description": "[virustotal] 8.8.8.8: Clean IP",
                        "detail": "info",
                        "source": "virustotal",
                        "severity": "info",
                        "summary": "Clean IP",
                    },
                    {
                        "event_type": "analysis",
                        "timestamp": "2026-07-06 10:02:00",
                        "description": "AI analysis by gemini-2.5-flash",
                        "model": "gemini-2.5-flash",
                        "summary": "This IP is benign. No threats detected.",
                    },
                ],
            },
            {
                "ioc": {
                    "id": 2,
                    "type": "domain",
                    "value": "evil.com",
                    "added_at": "2026-07-06 10:03:00",
                },
                "events": [
                    {
                        "event_type": "enrichment",
                        "timestamp": "2026-07-06 10:04:00",
                        "description": "[urlhaus] evil.com: Malware distribution",
                        "detail": "high",
                        "source": "urlhaus",
                        "severity": "high",
                        "summary": "Malware distribution",
                    },
                ],
            },
        ],
        "case_notes": [
            {
                "id": 1,
                "content": "Escalated to IR team.",
                "created_at": "2026-07-06 10:10:00",
            }
        ],
        "stats": {
            "ioc_count": 2,
            "enrichment_count": 2,
            "analysis_count": 1,
            "top_severity": "high",
        },
    }


class TestCaseExportDocx:
    """Test export_case_to_docx output."""

    def test_docx_produces_bytes(self, sample_history, tmp_path):
        """export_case_to_docx should write a .docx file and return path."""
        from lupe.export.docx_export import export_case_to_docx

        output_path = tmp_path / "report.docx"
        result = export_case_to_docx(sample_history, output_path)
        assert result.exists()
        assert result.suffix == ".docx"

    def test_docx_contains_case_name(self, sample_history, tmp_path):
        """DOCX should contain the case name in a heading."""
        from docx import Document

        from lupe.export.docx_export import export_case_to_docx

        output_path = tmp_path / "report.docx"
        result = export_case_to_docx(sample_history, output_path)
        doc = Document(str(result))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Incident-2026-001" in full_text

    def test_docx_contains_ioc_values(self, sample_history, tmp_path):
        """DOCX should reference all IOC values."""
        from docx import Document

        from lupe.export.docx_export import export_case_to_docx

        output_path = tmp_path / "report.docx"
        result = export_case_to_docx(sample_history, output_path)
        doc = Document(str(result))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "8.8.8.8" in full_text
        assert "evil.com" in full_text

    def test_docx_contains_enrichment_sources(self, sample_history, tmp_path):
        """DOCX should reference enrichment source names."""
        from docx import Document

        from lupe.export.docx_export import export_case_to_docx

        output_path = tmp_path / "report.docx"
        result = export_case_to_docx(sample_history, output_path)
        doc = Document(str(result))
        # Check tables for enrichment data
        table_text = ""
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    table_text += cell.text + " "
        assert "virustotal" in table_text
        assert "urlhaus" in table_text

    def test_docx_has_headings(self, sample_history, tmp_path):
        """DOCX should have proper heading structure."""
        from docx import Document

        from lupe.export.docx_export import export_case_to_docx

        output_path = tmp_path / "report.docx"
        result = export_case_to_docx(sample_history, output_path)
        doc = Document(str(result))
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        assert len(headings) >= 2  # At least case title + one IOC heading

    def test_docx_empty_case(self, tmp_path):
        """DOCX should handle empty case."""
        from lupe.export.docx_export import export_case_to_docx

        history = {
            "case": {
                "id": 1, "name": "Empty Case", "description": "",
                "status": "open", "created_at": "2026-01-01",
            },
            "iocs": [],
            "case_notes": [],
            "stats": {
                "ioc_count": 0, "enrichment_count": 0,
                "analysis_count": 0, "top_severity": "info",
            },
        }
        output_path = tmp_path / "empty.docx"
        result = export_case_to_docx(history, output_path)
        assert result.exists()


class TestIOCExportDocx:
    """Test export_ioc_to_docx output."""

    def test_single_ioc_docx(self, sample_history, tmp_path):
        """export_ioc_to_docx should produce a DOCX with only one IOC."""
        from docx import Document

        from lupe.export.docx_export import export_ioc_to_docx

        ioc_data = sample_history["iocs"][0]
        case_info = sample_history["case"]
        output_path = tmp_path / "ioc_report.docx"
        result = export_ioc_to_docx(ioc_data, case_info, output_path)
        assert result.exists()

        doc = Document(str(result))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "8.8.8.8" in full_text
